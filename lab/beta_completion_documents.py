"""Live synthetic local-worker format acceptance; no cloud pickup, emails or submissions."""
import io,json,hashlib,time
from pathlib import Path
from playwright.sync_api import sync_playwright
import fitz
from docx import Document
from openpyxl import Workbook
root=Path('/private');data=json.loads((root/'credentials.json').read_text());base=data['base'];state=json.loads((root/'state.json').read_text())
assert base=='https://grantflow-modern-staging.onrender.com' and data['prefix'].startswith('beta_verify_')
proof='SYNTHETIC BLUEBIRD EDUCATION VERIFICATION'
pdf=fitz.open();page=pdf.new_page(width=700,height=180);page.insert_text((30,70),proof,fontsize=22);png=page.get_pixmap().tobytes('png');scan=fitz.open();scan.new_page(width=700,height=180).insert_image(fitz.Rect(0,0,700,180),stream=png)
doc=Document();doc.add_paragraph(proof);docbuf=io.BytesIO();doc.save(docbuf)
book=Workbook();book.active.append(['SYNTHETIC BLUEBIRD EDUCATION','Amount']);book.active.append(['Synthetic revenue','100.00']);bookbuf=io.BytesIO();book.save(bookbuf)
files=[('md',proof.encode(),'text/markdown'),('pdf',pdf.tobytes(),'application/pdf'),('docx',docbuf.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),('csv',b'SYNTHETIC BLUEBIRD EDUCATION,Amount\nSynthetic revenue,100.00\n','text/csv'),('xlsx',bookbuf.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),('png',png,'image/png'),('scanned.pdf',scan.tobytes(),'application/pdf')];out=[]
with sync_playwright()as p:
 def session(role):
  a=data['accounts'][role];c=p.request.new_context();r=c.post(base+'/gfm-app/api/auth/login',headers={'Origin':base},data={'username':a['username'],'password':a['password']});assert r.ok;return c
 writer=session('grant_writer');admin=session('admin')
 def headers(c):return {'Origin':base,'X-CSRF-Token':next(x['value']for x in c.storage_state()['cookies']if x['name']=='__Host-gfm_csrf')}
 for suffix,raw,mime in files:
  started=time.monotonic();r=writer.post(base+'/gfm-app/api/ingestion/upload',headers=headers(writer),multipart={'file':{'name':'synthetic-format-'+data['nonce']+'-verified.'+suffix,'mimeType':mime,'buffer':raw},'application_id':str(state['application_id'])});assert r.status==202;item=r.json();did=item['id']
  until=time.monotonic()+90
  while time.monotonic()<until:
   item=writer.get(base+'/gfm-app/api/ingestion/documents/'+str(did)).json()
   if item['status'] in ('ready','needs_review','failed'):break
   time.sleep(1)
  if item['status']=='failed' and 'Local OCR requires at least 2048 MB' in (item.get('error') or ''):
   blob=writer.get(base+'/gfm-app/api/ingestion/documents/'+str(did)+'/original');assert blob.ok and hashlib.sha256(blob.body()).hexdigest()==hashlib.sha256(raw).hexdigest()
   out.append({'format':suffix,'result':'BLOCKED','reason':item['error'],'original_upload_preserved':True});print(json.dumps(out[-1]),flush=True);(root/'format-results.json').write_text(json.dumps(out,indent=2));continue
  assert item['status'] in ('ready','needs_review'),(suffix,item['status'],item.get('error'))
  reviewed=item['status']=='needs_review'
  if reviewed:
   body={'confirmed':True,'sha256':item['sha256'],'revision':item['extraction_revision']};url=base+'/gfm-app/api/ingestion/documents/'+str(did)+'/review'
   assert writer.post(url,headers=headers(writer),data=body).status==403
   result=admin.post(url,headers=headers(admin),data=body);assert result.ok and result.json()['status']=='ready'
  search=writer.get(base+'/gfm-app/api/ingestion/search?q=SYNTHETIC').json();assert any(x['document']['id']==did for x in search),(suffix,'Search did not include approved extracted text')
  blob=writer.get(base+'/gfm-app/api/ingestion/documents/'+str(did)+'/original');assert blob.ok and hashlib.sha256(blob.body()).hexdigest()==hashlib.sha256(raw).hexdigest()
  out.append({'format':suffix,'result':'PASS','actual_ocr':bool(item['ocr']),'review_required_and_approved':reviewed,'seconds':round(time.monotonic()-started,3)});print(json.dumps(out[-1]),flush=True);(root/'format-results.json').write_text(json.dumps(out,indent=2))
 writer.dispose();admin.dispose()
