"""Generate unmistakably synthetic financial documents for ingestion QA."""
import argparse
import json
from pathlib import Path
import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def create_samples(folder):
    folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
    pdf=pymupdf.open();page=pdf.new_page()
    lines=[('SYNTHETIC TEST DATA - NOT REAL FINANCIALS',18),
           ('Example Learning Foundation - fictional organization',13),
           ('Financial statement | Test period: October 2026',13),
           ('Revenue: USD 18432.75',18),('Expenses: USD 6001.25',18),
           ('Net surplus: USD 12431.50',18),
           ('Mission: improve education and science learning opportunities.',12),
           ('Created solely for GrantFlow ingestion testing. Not an audit or award.',11)]
    y=65
    for text,size in lines:
        page.insert_text((36,y),text,fontsize=size);y+=44
    pdf.save(folder/'SYNTHETIC-financial-statement.pdf')
    pix=page.get_pixmap(matrix=pymupdf.Matrix(2,2))
    png=folder/'SYNTHETIC-financial-scan.png';pix.save(png)
    scanned=pymupdf.open();scanpage=scanned.new_page(width=page.rect.width,height=page.rect.height)
    scanpage.insert_image(scanpage.rect,stream=pix.tobytes('png'))
    scanned.save(folder/'SYNTHETIC-scanned-statement.pdf');scanned.close();pdf.close()
    wb=Workbook();ws=wb.active;ws.title='Synthetic financials'
    for row in [('SYNTHETIC TEST DATA - NOT REAL FINANCIALS',),('Fictional organization','Example Learning Foundation'),
                ('Account','Amount USD'),('Revenue',18432.75),('Expenses',6001.25),('Net surplus',12431.50),
                ('Formula example - review required','=B4-B5')]:ws.append(row)
    for cell in ws[1]+ws[3]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='164E63')
    for row in range(4,8):ws.cell(row,2).number_format='#,##0.00'
    ws.column_dimensions['A'].width=55;ws.column_dimensions['B'].width=32
    wb.save(folder/'SYNTHETIC-financial-workbook.xlsx')
    (folder/'README.txt').write_text('All documents are fictional test data. Do not use in a real grant submission.\n'
        'Expected revenue: USD 18432.75; expenses: USD 6001.25; net surplus: USD 12431.50.\n'
        'The workbook includes a formula without a cached result; the importer must flag it, never calculate or invent its value.\n'
        'PDF/PNG scans require OCR. Financial extraction requires reviewer confirmation before AI use.\n',encoding='utf-8')
    return folder


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output',required=True)
    folder=create_samples(parser.parse_args().output)
    print(json.dumps({'folder':str(folder),'files':[p.name for p in folder.iterdir()]}))
