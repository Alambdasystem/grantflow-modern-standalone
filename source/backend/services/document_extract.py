"""Bounded, offline extraction. Never execute document macros, formulas or instructions."""
import csv
import io
import json
from pathlib import Path
import sys
import zipfile
import xml.etree.ElementTree as ET

EXTENSIONS = {'.pdf', '.docx', '.xlsx', '.csv', '.txt', '.md', '.png'}
MAX_BYTES = 20 * 1024 * 1024
MAX_TEXT = 500000
CATEGORIES = ('financials', 'mission', 'priorities', 'successful_grants', 'audits', 'award_history', 'general')


def ocr_table(result, locator):
    """Geometric reading rows, never inferred financial semantics or corrected amounts."""
    lines = []
    for polygon, text, confidence in result or []:
        xs, ys = [float(p[0]) for p in polygon], [float(p[1]) for p in polygon]
        lines.append({'text': str(text), 'confidence': float(confidence),
                      'bbox': [min(xs), min(ys), max(xs), max(ys)]})
    groups = []
    for cell in sorted(lines, key=lambda c: (c['bbox'][1], c['bbox'][0])):
        box = cell['bbox']
        center = (box[1] + box[3]) / 2
        match = next((g for g in groups if abs(g['center'] - center) <=
                      .4 * min(g['height'], max(1, box[3] - box[1]))), None)
        if match is None:
            match = {'center': center, 'height': max(1, box[3] - box[1]), 'cells': []}
            groups.append(match)
        match['cells'].append(cell)
    cells = [sorted(g['cells'], key=lambda c: c['bbox'][0]) for g in groups]
    return {'locator': locator, 'rows': [[c['text'] for c in row] for row in cells],
            'cells': cells, 'requires_review': True, 'layout': 'ocr_geometric_rows'}


def extract(raw: bytes, filename: str) -> dict:
    suffix = Path(filename).suffix.lower()
    if suffix not in EXTENSIONS or not raw or len(raw) > MAX_BYTES:
        raise ValueError('Use a nonempty PDF, PNG, DOCX, XLSX, CSV or text file under 20 MB.')
    parts, tables, warnings = [], [], []
    ocr = False
    metadata = {'format': suffix[1:], 'extractor_version': 'gfm-2'}
    try:
        if suffix in {'.docx', '.xlsx'}:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                infos = archive.infolist()
                if len(infos) > 5000 or sum(f.file_size for f in infos) > 50 * 1024 * 1024:
                    raise ValueError('Expanded document exceeds safe processing limits.')
                if any('vbaproject' in f.filename.lower() for f in infos):
                    raise ValueError('Macro-enabled content is not supported.')
        if suffix == '.pdf':
            import pymupdf
            with pymupdf.open(stream=raw, filetype='pdf') as pdf:
                if pdf.needs_pass or len(pdf) > 100:
                    raise ValueError('Use an unencrypted PDF with at most 100 pages.')
                metadata['pages'] = len(pdf)
                reader = None
                for page in pdf:
                    text = page.get_text(sort=True).strip()
                    locator = f'page {page.number + 1}'
                    if not text or (len(text)<100 and page.get_images()):
                        ocr = True
                        if reader is None:
                            from rapidocr_onnxruntime import RapidOCR
                            reader = RapidOCR(intra_op_num_threads=1, inter_op_num_threads=1)
                        import numpy as np
                        scale = min(2.0, (12000000 / max(1, page.rect.width * page.rect.height)) ** .5)
                        pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), colorspace=pymupdf.csRGB)
                        result, _ = reader(np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, 3))
                        text = '\n'.join(str(line[1]) for line in result or [])
                        tables.append(ocr_table(result, locator))
                        warnings.append(f'{locator}: OCR text requires review against the original, including all amounts and table structure.')
                    else:
                        for table in page.find_tables().tables:
                            tables.append({'locator': locator, 'rows': table.extract()})
                    parts.append({'locator': locator, 'text': text})
        elif suffix == '.png':
            from PIL import Image
            import numpy as np
            from rapidocr_onnxruntime import RapidOCR
            with Image.open(io.BytesIO(raw)) as image:
                if image.format != 'PNG' or image.width * image.height > 12000000:
                    raise ValueError('Use a valid PNG with at most 12 million pixels.')
                if getattr(image, 'n_frames', 1) != 1:
                    raise ValueError('Animated PNG documents are not supported.')
                metadata.update(width=image.width, height=image.height)
                rgba=image.convert('RGBA')
                background=Image.new('RGBA',rgba.size,'white')
                background.alpha_composite(rgba)
                reader=RapidOCR(intra_op_num_threads=1,inter_op_num_threads=1)
                result,_=reader(np.array(background.convert('RGB')))
                parts=[{'locator':'image 1','text':'\n'.join(str(line[1]) for line in result or [])}]
                tables.append(ocr_table(result, 'image 1'))
                ocr=True
                warnings.append('PNG OCR requires review against the original, including all amounts and table structure.')
        elif suffix == '.docx':
            from docx import Document
            doc = Document(io.BytesIO(raw))
            parts = [{'locator': f'paragraph {i+1}', 'text': p.text} for i,p in enumerate(doc.paragraphs) if p.text.strip()]
            for i, table in enumerate(doc.tables):
                rows = [[c.text for c in r.cells] for r in table.rows]
                tables.append({'locator': f'table {i+1}', 'rows': rows})
        elif suffix == '.xlsx':
            # Read literal OOXML values, not binary floats; never evaluate formulas or external links.
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                ns = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                strings = []
                if 'xl/sharedStrings.xml' in archive.namelist():
                    root = ET.fromstring(archive.read('xl/sharedStrings.xml'))
                    strings = [''.join(n.itertext()) for n in root.findall('s:si', ns)]
                sheets = sorted(n for n in archive.namelist() if n.startswith('xl/worksheets/sheet') and n.endswith('.xml'))
                if len(sheets) > 100:
                    raise ValueError('Spreadsheet exceeds 100 sheets.')
                for name in sheets:
                    root = ET.fromstring(archive.read(name))
                    rows = []
                    for row in root.findall('.//s:sheetData/s:row', ns):
                        values = []
                        for c in row.findall('s:c', ns):
                            value = c.findtext('s:v', default='', namespaces=ns)
                            if c.get('t') == 's': value = strings[int(value)]
                            if c.get('t') == 'inlineStr': value = ''.join(c.find('s:is', ns).itertext())
                            if c.find('s:f', ns) is not None:
                                warnings.append('Spreadsheet formulas are not evaluated; cached values require review.')
                                value = value or '[formula has no cached value]'
                            values.append(f'{c.get("r")}: {value}')
                        rows.append(values)
                        if len(rows) > 20000: raise ValueError('Spreadsheet exceeds 20,000 rows per sheet.')
                    tables.append({'locator': name, 'rows': rows})
                warnings.append('Spreadsheet values preserve stored numeric text; currency/date formatting must be checked in the original.')
        elif suffix == '.csv':
            rows = list(csv.reader(io.StringIO(raw.decode('utf-8-sig')), strict=True))
            if len(rows) > 20000: raise ValueError('CSV exceeds 20,000 rows.')
            tables = [{'locator': 'CSV rows', 'rows': rows}]
        else:
            parts = [{'locator':'text', 'text':raw.decode('utf-8-sig')}]
        for table in tables:
            parts.append({'locator':table['locator'], 'text':'\n'.join(' | '.join(str(c or '') for c in row) for row in table['rows'])})
        text = '\n\n'.join(f'[{p["locator"]}]\n{p["text"]}' for p in parts)
        if not any(p['text'].strip() for p in parts) or len(text)>MAX_TEXT:
            raise ValueError('No readable text, or extracted content exceeds 500,000 characters.')
        haystack = (filename + ' ' + text[:10000]).lower()
        category = 'general'
        for key, terms in [('audits',('audit report','auditor')),('award_history',('award history',)),('successful_grants',('successful grant',)),('financials',('financial','balance sheet','income statement','cash flow','liability','revenue')),('mission',('mission',)),('priorities',('priorities','strategic plan'))]:
            if any(term in haystack for term in terms):
                category = key
                break
        if category == 'financials': warnings.append('Financial values require human verification; no amounts are normalized, rounded or inferred.')
        return {'text':text, 'parts':parts, 'tables':tables, 'metadata':metadata, 'category':category,
                'warnings':list(dict.fromkeys(warnings)), 'needs_review':bool(warnings), 'ocr':ocr}
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Document is malformed, encrypted, or could not be extracted. Review the source and retry.') from exc


if __name__ == '__main__':
    try:
        # Parser diagnostics must never corrupt the worker's JSON protocol.
        from contextlib import redirect_stdout
        with redirect_stdout(sys.stderr):
            result = extract(Path(sys.argv[1]).read_bytes(), sys.argv[2])
        print(json.dumps(result, ensure_ascii=True))
    except Exception as exc:
        print(json.dumps({'error': str(exc) if isinstance(exc, ValueError) else 'Extraction failed.'}))
        sys.exit(1)
