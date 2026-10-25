"""Document extraction and tenant-bound ingestion regression tests."""
import sys
from pathlib import Path
import pytest
import asyncio
import io
import json

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from test_identity_ownership import client, headers


@pytest.fixture
def api(client):
    from routers import ingestion, operations, ai_router
    for router in (ingestion.router,operations.router,ai_router.router): client.app.include_router(router)
    return client


def upload(api,name='mission.txt',data=b'Mission: synthetic education learning evidence.',category='general',role='writer',**kwargs):
    r=api.post('/api/ingestion/upload',files={'file':(name,data)},data={'category':category,**kwargs},headers=headers(role))
    assert r.status_code==202,r.text
    return r.json()


def drain():
    from db import database
    from models.models import AutomationJob,AutomationJobStatus
    from sqlalchemy import select
    from services.document_ingestion import process
    async def run():
        async with database.AsyncSessionLocal() as db:
            ids=(await db.execute(select(AutomationJob.id).where(AutomationJob.job_type=='ingest',AutomationJob.status==AutomationJobStatus.queued))).scalars().all()
        for id in ids: await process(id)
    asyncio.run(run())


def test_upload_index_duplicates_versions_and_permissions(api):
    item=upload(api);drain()
    result=api.get('/api/ingestion/search?q=education',headers=headers()).json()
    assert len(result)==1 and 'document 1 v1' in result[0]['citation']
    assert upload(api)['duplicate'] is True
    assert len(api.get('/api/ingestion/documents',headers=headers()).json())==1
    v2=upload(api,data=b'Mission: updated science evidence.');assert v2['version']==2
    assert api.get('/api/ingestion/search?q=education',headers=headers()).json()==[]
    drain()
    assert api.get('/api/ingestion/search?q=science',headers=headers()).json()
    assert api.get('/api/ingestion/documents').status_code==401
    assert api.get('/api/ingestion/documents/1/original').status_code==401
    assert api.post('/api/ingestion/upload',files={'file':('a.txt',b'text')},headers=headers('viewer')).status_code==403
    assert api.post('/api/ingestion/upload',files={'file':('a.txt',b'text')},data={'application_id':2},headers=headers()).status_code==404
    assert api.get('/api/ingestion/documents/999',headers=headers('admin')).status_code==404


def test_review_gate_and_provenance_deletion(api,monkeypatch):
    from routers import ai_router
    from unittest.mock import AsyncMock
    item=upload(api,'financial.csv',b'Account,Amount\nRevenue,1234.50\n');drain()
    assert api.get(f'/api/ingestion/documents/{item["id"]}',headers=headers()).json()['status']=='needs_review'
    assert api.get('/api/ingestion/search?q=1234.50',headers=headers()).json()==[]
    body={'confirmed':True,'sha256':item['sha256']}
    assert api.post(f'/api/ingestion/documents/{item["id"]}/review',json=body,headers=headers()).status_code==403
    assert api.post(f'/api/ingestion/documents/{item["id"]}/review',json=body,headers=headers('admin')).status_code==200
    monkeypatch.setattr(ai_router._ai,'generate',AsyncMock(return_value='Synthetic source-grounded draft'))
    r=api.post('/api/operations/drafts',json={'application_id':1,'task':'loi'},headers=headers())
    assert r.status_code==201,r.text
    draft_id=r.json()['id']
    monkeypatch.setattr(ai_router._ai,'chat',AsyncMock(return_value='Revenue 1234.50'))
    chat=api.post('/api/ai/chat',json={'stream':False,'messages':[{'role':'user','content':'Revenue?'}]},headers=headers())
    assert chat.status_code==200,chat.text
    assert 'financials' not in api.get('/api/operations/knowledge-status',headers=headers()).json()['missing_categories']
    assert api.delete(f'/api/ingestion/documents/{item["id"]}',headers=headers('admin')).status_code==200
    assert api.get('/api/ingestion/search?q=1234.50',headers=headers()).json()==[]
    assert 'Source document deleted' in api.get(f'/api/operations/records/{draft_id}/export',headers=headers()).text
    assert api.get('/api/ai/conversations',headers=headers()).json()[0]['messages']==[]
    assert api.get(f'/api/ingestion/documents/{item["id"]}/original',headers=headers()).status_code==410
    assert upload(api,'financial.csv',b'Account,Amount\nRevenue,1234.50\n')['status']=='deleted'


def test_failed_parser_retry_and_other_tenant(api):
    from db import database
    from services.document_ingestion import receive
    bad=upload(api,'broken.pdf',b'not PDF');drain()
    detail=api.get(f'/api/ingestion/documents/{bad["id"]}',headers=headers()).json()
    assert detail['status']=='failed' and detail['retry_at'] and detail['history'][-1]['action']=='failed'
    assert api.post(f'/api/ingestion/documents/{bad["id"]}/retry',headers=headers()).status_code==200
    async def other():
        async with database.AsyncSessionLocal() as db:
            row,_=await receive(db,2,2,b'Secret other organization','private.txt');return row.id
    id=asyncio.run(other());drain()
    assert api.get(f'/api/ingestion/documents/{id}',headers=headers('admin')).status_code==404
    assert api.get(f'/api/ingestion/documents/{id}/original',headers=headers('admin')).status_code==404
    assert api.delete(f'/api/ingestion/documents/{id}',headers=headers('admin')).status_code==404
    assert api.get('/api/ingestion/search?q=Secret',headers=headers('admin')).json()==[]


def test_scheduled_pickup_and_source_outage(api,tmp_path,monkeypatch):
    from services import document_ingestion as service
    from config import settings
    from db import database
    from models.models import IngestionSource
    import os,time
    folder=tmp_path/'authorized-test-source';folder.mkdir()
    path=folder/'mission.csv';path.write_text('Mission,education\nPriority,science',encoding='utf-8');os.utime(path,(time.time()-10,)*2)
    manifest=tmp_path/'sources.json';manifest.write_text(json.dumps([{'key':'test','label':'Authorized synthetic test','path':str(folder),'organization_id':1,'created_by_id':1,'enabled':True,'interval_seconds':60}]))
    monkeypatch.setattr(settings,'INGESTION_SOURCE_MANIFEST',str(manifest))
    asyncio.run(service.schedule());drain()
    assert api.get('/api/ingestion/search?q=science',headers=headers()).json()
    assert api.get('/api/ingestion/sources',headers=headers()).json()['sources'][0]['last_success']
    folder.rename(tmp_path/'temporarily-offline')
    async def due():
        async with database.AsyncSessionLocal() as db:
            s=await db.get(IngestionSource,1);s.next_poll=None;await db.commit()
    asyncio.run(due());asyncio.run(service.schedule())
    assert api.get('/api/ingestion/sources',headers=headers()).json()['sources'][0]['error']
    assert api.get('/api/ingestion/search?q=science',headers=headers()).json()
    (tmp_path/'temporarily-offline').rename(folder)
    asyncio.run(due());asyncio.run(service.schedule())
    assert api.get('/api/ingestion/sources',headers=headers()).json()['sources'][0]['error'] is None


def test_docx_xlsx_and_pdf():
    from services.document_extract import extract
    from docx import Document
    from openpyxl import Workbook
    import pymupdf
    doc=Document();doc.add_paragraph('Mission: education');doc.add_table(rows=1,cols=2).rows[0].cells[0].text='1234.50'
    buf=io.BytesIO();doc.save(buf)
    assert '1234.50' in extract(buf.getvalue(),'mission.docx')['text']
    wb=Workbook();ws=wb.active;ws.append(['Revenue','1234.50']);ws['B2']='=1+2';buf=io.BytesIO();wb.save(buf)
    result=extract(buf.getvalue(),'financial.xlsx');assert '1234.50' in result['text'] and result['needs_review']
    pdf=pymupdf.open();page=pdf.new_page();page.insert_text((72,72),'Mission: education');raw=pdf.tobytes();pdf.close()
    assert 'education' in extract(raw,'mission.pdf')['text']


def test_scanned_pdf_ocr():
    from services.document_extract import extract
    import pymupdf
    pdf=pymupdf.open();p=pdf.new_page();p.insert_text((72,72),'Revenue 1234.50',fontsize=24);pix=p.get_pixmap(matrix=pymupdf.Matrix(2,2))
    scan=pymupdf.open();page=scan.new_page();page.insert_image(page.rect,stream=pix.tobytes('png'))
    result=extract(scan.tobytes(),'financial-scan.pdf')
    assert result['ocr'] and result['needs_review'] and '1234.50' in result['text']


def test_png_ocr_and_invalid_image():
    from services.document_extract import extract
    import pymupdf
    pdf=pymupdf.open();page=pdf.new_page();page.insert_text((72,72),'SYNTHETIC TEST Revenue 1234.50',fontsize=20)
    raw=page.get_pixmap(matrix=pymupdf.Matrix(2,2)).tobytes('png')
    result=extract(raw,'financial-test.png')
    assert result['ocr'] and result['needs_review'] and '1234.50' in result['text']
    assert result['metadata']['width']>0
    with pytest.raises(ValueError):extract(b'not an image','bad.png')


def test_pdf_worker_protocol_excludes_parser_diagnostics(tmp_path):
    import subprocess
    import pymupdf
    from services import document_extract
    pdf=pymupdf.open();page=pdf.new_page();page.insert_text((72,72),'SYNTHETIC Revenue 18432.75')
    path=tmp_path/'financial.pdf';pdf.save(path);pdf.close()
    response=subprocess.run([sys.executable,document_extract.__file__,str(path),path.name],capture_output=True,check=True)
    result=json.loads(response.stdout)
    assert '18432.75' in result['text'] and result['needs_review']


def test_csv_preserves_financial_lexemes():
    from services.document_extract import extract
    result = extract(b'Account,Amount\nCash,"1,234.50"\nLiability,(400.00)\n', 'financial.csv')
    assert result['tables'][0]['rows'][1] == ['Cash', '1,234.50']
    assert '(400.00)' in result['text']
    assert result['category'] == 'financials'
    assert result['needs_review'] is True


def test_malformed_and_unsupported():
    from services.document_extract import extract
    for raw, name in [(b'not a PDF','bad.pdf'),(b'bad zip','bad.docx'),(b'run','bad.exe')]:
        with pytest.raises(ValueError):
            extract(raw, name)
