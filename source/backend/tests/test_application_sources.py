"""Disposable document-to-application integration tests; provider fixtures are labelled."""
import asyncio
import hashlib
import io
import json
import zipfile
from unittest.mock import AsyncMock
import pytest
from test_identity_ownership import client,headers
from db import database
from models.models import IngestionDocument,Document,GrantApplication,ApplicationSourceRun as Run
from services import application_sources as service

@pytest.fixture
def api(client,monkeypatch):
    from config import settings
    monkeypatch.setattr(settings,'AI_ENABLED',True)
    from routers.application_sources import router
    client.app.include_router(router)
    async def seed():
        async with database.AsyncSessionLocal() as db:
            app=await db.get(GrantApplication,1);app.amount_requested=100000
            app.form_data={'honda_2027_onboarding':{'total_program_budget':425000,'other_funding_needed_not_secured':325000}}
            from config import settings
            from pathlib import Path
            p=Path(settings.INGESTION_STORAGE_DIR);p.mkdir(parents=True,exist_ok=True)
            raw=b'Please describe your program. Under 2,000 Characters';f=p/'source.txt';f.write_bytes(raw)
            db.add(Document(id=1,application_id=1,organization_id=1,filename='source.txt',original_filename='source.txt',file_path=str(f),file_size=len(raw)))
            await db.flush()
            db.add(IngestionDocument(id=1,organization_id=1,created_by_id=1,document_id=1,logical_key='fixture',sha256=hashlib.sha256(raw).hexdigest(),version=1,status='needs_review',current=True,
                extracted={'text':raw.decode(),'parts':[{'locator':'paragraph 1','text':raw.decode()}],'warnings':['Unverified draft']}))
            await db.commit()
    asyncio.run(seed())
    return client

def inventory(api,monkeypatch,authorities=None):
    raw=json.dumps({'items':[{'kind':'question','wording':'Please describe your program.','source_id':1,'locator':'paragraph 1','source_quote':'Please describe your program. Under 2,000 Characters','required':True,'human_supplied':False}]})
    monkeypatch.setattr(service.AIService,'chat',AsyncMock(return_value=raw))
    r=api.post('/api/applications/1/sources/analyze',json={'source_ids':[1],'authorities':authorities or {}},headers=headers());assert r.status_code==202,r.text
    id=r.json()['id'];asyncio.run(service.process(id))
    rows=api.get('/api/applications/1/sources',headers=headers()).json()['runs'];return next(r for r in rows if r['id']==id)

def test_inventory_provenance_dedup_unverified_and_main_preserved(api,monkeypatch):
    row=inventory(api,monkeypatch);assert row['status']=='completed'
    item=row['output']['items'][0]
    assert item['authority']=='supplied_outline' and item['verification']=='needs_review' and item['unverified_draft']
    assert item['limit']=={'value':2000,'unit':'characters','operator':'lt'}
    transport={k:item[k] for k in ('kind','wording','source_id','locator','source_quote','required','human_supplied')}
    batch=[{'source_id':1,'locator':'paragraph 1','text':item['source_quote']}]
    for raw in [json.dumps([transport]),json.dumps([{'items':[transport]}])]:
        assert service.parse_items(raw,batch,row['sources'])[0]['id']==item['id']
    again=api.post('/api/applications/1/sources/analyze',json={'source_ids':[1]},headers=headers()).json();assert again['id']==row['id']
    assert not api.get('/api/applications/1',headers=headers()).json()['proposal_text']
    assert api.get('/api/applications/2/sources',headers=headers()).status_code==404
    assert api.post('/api/applications/1/sources/analyze',json={'source_ids':[1]},headers=headers('viewer')).status_code==403
    assert api.post('/api/applications/2/sources/analyze',json={'source_ids':[1]},headers=headers('admin')).status_code==404

def test_fabricated_questions_and_citations_are_withheld(api,monkeypatch):
    monkeypatch.setattr(service.AIService,'chat',AsyncMock(return_value=json.dumps({'items':[{'kind':'question','wording':'Invented official question','source_id':1,'locator':'paragraph 1','source_quote':'Please describe your program.','required':True,'human_supplied':False}]})))
    r=api.post('/api/applications/1/sources/analyze',json={'source_ids':[1]},headers=headers()).json();asyncio.run(service.process(r['id']))
    run=api.get('/api/applications/1/sources',headers=headers()).json()['runs'][0]
    assert run['status']=='failed' and 'Inferred question' in run['error']
    assert not run['output'].get('items') and not run['output']['checkpoint']['items']
    assert len(run['output']['calls'])==2 and run['output']['checkpoint']['withheld']
    assert api.post(f'/api/applications/1/sources/runs/{r["id"]}/retry',headers=headers()).status_code==202

def test_checkpoint_retry_keeps_valid_items_and_does_not_repeat_completed_batch(api,monkeypatch):
    raw=json.dumps({'items':[{'kind':'question','wording':'Please describe your program.','source_id':1,'locator':'paragraph 1','source_quote':'Please describe your program.','required':None,'human_supplied':False}]})
    real_batches=service.batches
    monkeypatch.setattr(service,'batches',lambda parts:[real_batches(parts)[0],real_batches(parts)[0]])
    chat=AsyncMock(side_effect=[raw,RuntimeError('private provider detail')])
    monkeypatch.setattr(service.AIService,'chat',chat)
    row=api.post('/api/applications/1/sources/analyze',json={'source_ids':[1]},headers=headers()).json()
    asyncio.run(service.process(row['id']))
    failed=api.get('/api/applications/1/sources',headers=headers()).json()['runs'][0]
    assert failed['status']=='failed'
    assert failed['output']['checkpoint']['completed_batches']==[0]
    assert len(failed['output']['checkpoint']['items'])==1
    assert failed['output']['calls'][-1]['status']=='failed'
    assert 'private provider detail' not in failed['error']
    assert api.post(f'/api/applications/1/sources/runs/{row["id"]}/retry',headers=headers()).status_code==202
    chat=AsyncMock(return_value='{"items":[]}')
    monkeypatch.setattr(service.AIService,'chat',chat)
    asyncio.run(service.process(row['id']))
    final=api.get('/api/applications/1/sources',headers=headers()).json()['runs'][0]
    assert final['status']=='completed' and final['attempts']==2
    assert len(final['output']['items'])==1 and not final['output']['inventory_complete']
    assert chat.await_count==1

def test_bounded_draft_context_retains_budget_table_and_reports_omissions():
    parts=[{'source_id':22,'locator':f'paragraph {i}','text':'unrelated '*600} for i in range(20)]
    parts += [{'source_id':23,'locator':'xl/worksheets/sheet1.xml','text':'Personnel $30,000; training $6,000.'},
              {'source_id':22,'locator':'paragraph 25','text':'Budget question'}]
    selected,coverage=service.draft_context(parts,{'source_id':22,'locator':'paragraph 25','wording':'Budget question'},'budget_narrative')
    assert coverage['selected_characters']<=4800 and coverage['omitted']
    assert not coverage['complete_source_coverage']
    assert any(p['source_id']==23 for p in selected)
    assert any(p['locator']=='paragraph 25' for p in selected)

def test_validation_repair_retains_failed_candidate_and_does_not_approve(api,monkeypatch):
    inv=inventory(api,monkeypatch)
    chat=AsyncMock(side_effect=['Invented 999 participants.',
        'Proposed pilot [information needed: dates] [source 1, paragraph 1]'])
    monkeypatch.setattr(service.AIService,'chat',chat)
    row=api.post('/api/applications/1/sources/draft',json={'inventory_id':inv['id'],'requirement_id':inv['output']['items'][0]['id']},headers=headers()).json()
    asyncio.run(service.process(row['id']))
    draft=api.get('/api/applications/1/sources',headers=headers()).json()['runs'][0]
    assert draft['status']=='completed' and len(draft['output']['generation_candidates'])==2
    assert draft['output']['generation_candidates'][0]['body']=='Invented 999 participants.'
    assert draft['output']['validation']==['Missing information placeholders remain.']
    assert not draft['reviews'] and chat.await_count==2
    decision={'item_key':'draft','decision':'accepted','notes':'Fixture validation check','output_hash':draft['output_hash']}
    assert api.post(f'/api/applications/1/sources/runs/{draft["id"]}/review',json=decision,headers=headers('admin')).status_code==422
    revised=api.post(f'/api/applications/1/sources/runs/{draft["id"]}/revision',json={'body':'Proposed pilot [source 1, paragraph 1]','output_hash':draft['output_hash']},headers=headers()).json()
    assert revised['validation']==[]
    final=api.get('/api/applications/1/sources',headers=headers()).json()['runs'][0]
    decision['output_hash']=final['output_hash']
    assert api.post(f'/api/applications/1/sources/runs/{final["id"]}/review',json=decision,headers=headers('admin')).status_code==201

def test_review_bundle_includes_labelled_drafts_without_approval(api,monkeypatch):
    inv=inventory(api,monkeypatch)
    monkeypatch.setattr(service.AIService,'chat',AsyncMock(return_value='Proposed pilot [information needed: dates] [source 1, paragraph 1]'))
    row=api.post('/api/applications/1/sources/draft',json={'inventory_id':inv['id'],'requirement_id':inv['output']['items'][0]['id']},headers=headers()).json()
    asyncio.run(service.process(row['id']))
    normal=api.get('/api/applications/1/sources/export',headers=headers())
    with zipfile.ZipFile(io.BytesIO(normal.content)) as z:assert not any(n.startswith('unapproved/') for n in z.namelist())
    response=api.get('/api/applications/1/sources/export?include_unapproved=true',headers=headers())
    with zipfile.ZipFile(io.BytesIO(response.content)) as z:
        manifest=json.loads(z.read('manifest.json'))
        assert manifest['package_status']=='INCOMPLETE' and manifest['include_unapproved']
        name=next(n for n in z.namelist() if n.startswith('unapproved/') and n.endswith('.docx'))
        from docx import Document as WordDocument
        text='\n'.join(p.text for p in WordDocument(io.BytesIO(z.read(name))).paragraphs)
        assert 'UNAPPROVED' in text and 'information needed: dates' in text
        assert 'VALIDATION.json' in z.namelist() and 'PORTAL-ENTRY.txt' in z.namelist()
    runs=api.get('/api/applications/1/sources',headers=headers()).json()['runs']
    draft=next(r for r in runs if r['id']==row['id'])
    assert not draft['reviews'] and 'Missing information placeholders remain.' in draft['output']['validation']
    # A newer inventory must not hide existing work or imply it answers new questions.
    newer=inventory(api,monkeypatch,authorities={'1':'official_guidance'})
    assert newer['id']!=inv['id']
    response=api.get('/api/applications/1/sources/export?include_unapproved=true',headers=headers())
    with zipfile.ZipFile(io.BytesIO(response.content)) as z:
        manifest=json.loads(z.read('manifest.json'))
        assert any(d['inventory_id']==inv['id'] and d['run_id']==row['id'] for d in manifest['review_drafts'])
        assert manifest['package_status']=='INCOMPLETE'
    async def withhold():
        async with database.AsyncSessionLocal() as db:
            r=await db.get(Run,newer['id'])
            r.output={**r.output,'withheld_batches':[{'batch':0,'reason':'Unsupported source quotation'}]}
            await db.commit()
    asyncio.run(withhold())
    newer=next(r for r in api.get('/api/applications/1/sources',headers=headers()).json()['runs'] if r['id']==newer['id'])
    response=api.post(f'/api/applications/1/sources/runs/{newer["id"]}/review',
        json={'item_key':'inventory_complete','decision':'accepted','notes':'Fixture checks cannot override withheld batches','source_url':'https://example.test/form','output_hash':newer['output_hash']},headers=headers('admin'))
    assert response.status_code==422 and 'Withheld source batches' in response.json()['detail']
    recovery=api.post(f'/api/applications/1/sources/runs/{newer["id"]}/retry',headers=headers())
    assert recovery.status_code==202 and recovery.json()['id']!=newer['id']
    old=next(r for r in api.get('/api/applications/1/sources',headers=headers()).json()['runs'] if r['id']==newer['id'])
    assert old['output']['withheld_batches'] and old['status']=='completed'

def test_review_cannot_approve_held_source_or_bypass_role(api,monkeypatch):
    run=inventory(api,monkeypatch);item=run['output']['items'][0]
    data={'item_key':item['id'],'decision':'accepted','notes':'Confirmed against original file','output_hash':run['output_hash']}
    assert api.post(f'/api/applications/1/sources/runs/{run["id"]}/review',json=data,headers=headers()).status_code==403
    assert api.post(f'/api/applications/1/sources/runs/{run["id"]}/review',json=data,headers=headers('admin')).status_code==422
    data['decision']='needs_information'
    assert api.post(f'/api/applications/1/sources/runs/{run["id"]}/review',json=data,headers=headers('admin')).status_code==201
    hist=api.get(f'/api/applications/1/sources/runs/{run["id"]}/reviews',headers=headers()).json();assert hist[0]['actor_id']==1

def test_real_path_fixture_draft_revision_review_export(api,monkeypatch,tmp_path):
    inv=inventory(api,monkeypatch);req=inv['output']['items'][0]
    monkeypatch.setattr(service.AIService,'chat',AsyncMock(return_value='Proposed program narrative [source 1, paragraph 1].'))
    payload={'inventory_id':inv['id'],'requirement_id':req['id'],'document_type':'work_plan'}
    r=api.post('/api/applications/1/sources/draft',json=payload,headers=headers());assert r.status_code==202,r.text
    id=r.json()['id'];asyncio.run(service.process(id))
    run=api.get('/api/applications/1/sources',headers=headers()).json()['runs'][0]
    assert run['status']=='completed' and run['output']['validation']==[]
    decision={'item_key':'draft','decision':'accepted','notes':'Synthetic test review only','output_hash':run['output_hash']}
    assert api.post(f'/api/applications/1/sources/runs/{id}/review',json=decision,headers=headers('admin')).status_code==201
    edit=api.post(f'/api/applications/1/sources/runs/{id}/revision',json={'body':'Changed draft [information needed] [source 1, paragraph 1]','output_hash':run['output_hash']},headers=headers())
    assert edit.status_code==201 and edit.json()['id']!=id
    exported=api.get('/api/applications/1/sources/export',headers=headers());assert exported.status_code==200
    path=tmp_path/'reopened.zip';path.write_bytes(exported.content)
    with zipfile.ZipFile(path) as z:
        manifest=json.loads(z.read('manifest.json'));assert manifest['package_status']=='INCOMPLETE' and not manifest['submission_performed']
        texts=[n for n in z.namelist() if n.startswith('answers/') and n.endswith('.txt')];assert texts and b'Proposed program narrative' in z.read(texts[0])
        docs=[n for n in z.namelist() if n.endswith('.docx')]
        from docx import Document as WordDocument
        assert 'Proposed program narrative' in '\n'.join(p.text for p in WordDocument(io.BytesIO(z.read(docs[0]))).paragraphs)
    assert not api.get('/api/applications/1',headers=headers()).json()['proposal_text']

def test_recovery_is_persistent_and_requires_explicit_retry(api,monkeypatch):
    row=inventory(api,monkeypatch)
    async def interrupt():
        async with database.AsyncSessionLocal() as db:
            r=await db.get(Run,row['id']);r.status='running';await db.commit()
    asyncio.run(interrupt());asyncio.run(service.recover())
    run=api.get('/api/applications/1/sources',headers=headers()).json()['runs'][0]
    assert run['status']=='failed' and 'restart' in run['error'] and run['output']['items']
    assert api.post(f'/api/applications/1/sources/runs/{run["id"]}/retry',headers=headers()).status_code==202
    async def exhausted():
        async with database.AsyncSessionLocal() as db:
            r=await db.get(Run,run['id']);r.status='failed';r.attempts=3;await db.commit()
    asyncio.run(exhausted())
    path=f'/api/applications/1/sources/runs/{run["id"]}/retry';reason={'reason':'Synthetic provider incident investigated and resolved'}
    assert api.post(path,headers=headers(),json=reason).status_code==409
    recovery=api.post(path,headers=headers('admin'),json=reason);assert recovery.status_code==202
    assert recovery.json()['id']!=run['id']
    assert api.post(path,headers=headers('admin'),json=reason).json()['id']==recovery.json()['id']

@pytest.mark.parametrize('count,invalid',[(1999,False),(2000,True),(2001,True)])
def test_strict_under_limit(count,invalid):
    output={'body':'x'*count,'limit':{'value':2000,'unit':'characters','operator':'lt'}}
    issues=service.draft_validation(output)
    assert any('required lt' in i for i in issues)==invalid

def test_budget_scopes_and_source_coverage():
    good={'body':'Overall program total: $425,000\nHonda request: $100,000\nUnmet funding gap: $325,000\n[source 1, paragraph 1]',
        'budget_scopes':{'requested':100000,'program_total':425000,'unmet_funding_not_committed':325000},
        'source_references':[{'source_id':1,'locator':'paragraph 1'}]}
    assert service.draft_validation(good)==[]
    assert service.draft_validation({**good,'body':good['body'].replace('$100,000','$110,000')})
    assert service.draft_validation({**good,'body':good['body'].replace('[source 1','[source 2')})
    mixed='Requested amount of $100,000 for the overall program totaling $425,000, with $325,000 in unmet funding. [source 1, paragraph 1]'
    assert service.draft_validation({**good,'body':mixed})==[]
    assert service.draft_validation({**good,'body':mixed.replace('$100,000','$110,000')})

def test_docx_recovers_hyperlinks_and_preserves_tables():
    from docx import Document as WordDocument
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from services.document_extract import extract
    doc=WordDocument();p=doc.add_paragraph('Website: ')
    link=OxmlElement('w:hyperlink');run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='https://example.test/grant';run.append(text);link.append(run);p._p.append(link)
    table=doc.add_table(rows=1,cols=1);table.cell(0,0).text='Budget $100000'
    raw=io.BytesIO();doc.save(raw);result=extract(raw.getvalue(),'outline.docx')
    assert 'https://example.test/grant' in result['text'] and result['tables'][0]['rows']==[['Budget $100000']]

def test_context_changes_invalidate_review_and_package(api,monkeypatch):
    run=inventory(api,monkeypatch)
    async def change():
        async with database.AsyncSessionLocal() as db:
            app=await db.get(GrantApplication,1);app.amount_requested=90000;await db.commit()
    asyncio.run(change())
    state=api.get('/api/applications/1/sources',headers=headers()).json()
    assert state['runs'][0]['stale']
    assert api.post('/api/applications/1/sources/draft',json={'inventory_id':run['id'],'requirement_id':run['output']['items'][0]['id']},headers=headers()).status_code==409

def test_auto_analysis_deduplicates_and_preserves_review_hold(api):
    async def run():
        async with database.AsyncSessionLocal() as db:
            one=await service.auto_enqueue(db,1,1);two=await service.auto_enqueue(db,1,1)
            assert one.id==two.id
            assert (await db.get(IngestionDocument,1)).status=='needs_review'
    asyncio.run(run())

def test_human_evidence_cannot_be_manufactured():
    assert not service.draftable({'kind':'upload','human_supplied':True,'wording':'Signed financial statement'})
    assert service.draftable({'kind':'upload','human_supplied':False,'wording':'Upload a work plan narrative'})
    assert '..' not in service.filename('a/../../outside.txt')
    assert '\\' not in service.filename('a\\..\\..\\outside.txt')
    assert service.profile_matches({'mission':'Empower young people'}, {'mission'})==[{'field':'mission','value':'Empower young people'}]

def test_real_provider_participant_scope_regression():
    # Observed with the real provider: it changed a total cohort into a per-workshop count.
    output={'body':'Four workshops, each for twenty students. [source ID: 3, locator: text]',
        'source_references':[{'source_id':3,'locator':'text'}],
        'count_scopes':[{'value':'twenty','unit':'students','scope':'total'}]}
    assert any('in total' in issue for issue in service.draft_validation(output))
    assert service.draft_validation({**output,'body':'Four workshops for twenty students in total. [source ID: 3, locator: text]'})==[]

@pytest.mark.parametrize('bad,corrected,issue',[
    ('A synthetic program narrative.', 'A synthetic program narrative. [source 1, paragraph 1]', 'No source references'),
    ('Each workshop serves twenty students. [source 1, paragraph 1]', 'Workshops serve twenty students in total. [source 1, paragraph 1]', 'in total'),
    ('Delivery costs: [information needed]. [source 1, paragraph 1]', 'Synthetic delivery costs: $100,000. [source 1, paragraph 1]', 'placeholders'),
])
def test_observed_ai_errors_block_approval_and_corrected_versions_pass(api,monkeypatch,bad,corrected,issue):
    # Labelled synthetic evidence, not invented Honda costs or portal verification.
    async def evidence():
        async with database.AsyncSessionLocal() as db:
            source=await db.get(IngestionDocument,1)
            source.extracted={'text':'Please describe your program. Under 2,000 Characters. Synthetic workshops serve twenty students in total. Synthetic delivery costs: $100,000.',
                'parts':[{'locator':'paragraph 1','text':'Please describe your program. Under 2,000 Characters. Synthetic workshops serve twenty students in total. Synthetic delivery costs: $100,000.'}]}
            await db.commit()
    asyncio.run(evidence())
    inv=inventory(api,monkeypatch);req=inv['output']['items'][0]
    monkeypatch.setattr(service.AIService,'chat',AsyncMock(return_value=bad))
    created=api.post('/api/applications/1/sources/draft',headers=headers(),json={'inventory_id':inv['id'],'requirement_id':req['id']}).json()
    asyncio.run(service.process(created['id']))
    run=next(r for r in api.get('/api/applications/1/sources',headers=headers()).json()['runs'] if r['id']==created['id'])
    assert any(issue in problem for problem in run['output']['validation'])
    decision={'item_key':'draft','decision':'accepted','notes':'Synthetic reviewer regression','output_hash':run['output_hash']}
    assert api.post(f'/api/applications/1/sources/runs/{run["id"]}/review',headers=headers('admin'),json=decision).status_code==422
    revision=api.post(f'/api/applications/1/sources/runs/{run["id"]}/revision',headers=headers(),json={'body':corrected,'output_hash':run['output_hash']})
    assert revision.status_code==201
    child=next(r for r in api.get('/api/applications/1/sources',headers=headers()).json()['runs'] if r['id']==revision.json()['id'])
    assert child['output']['validation']==[]
    decision['output_hash']=child['output_hash']
    assert api.post(f'/api/applications/1/sources/runs/{child["id"]}/review',headers=headers('admin'),json=decision).status_code==201
    original=next(r for r in api.get('/api/applications/1/sources',headers=headers()).json()['runs'] if r['id']==run['id'])
    assert original['output']['body']==bad and original['reviews'].get('draft',{}).get('decision')!='accepted'

def test_reviewed_package_requires_explicit_source_portal_and_answer_reviews(api,monkeypatch):
    from routers.ingestion import router
    api.app.include_router(router)
    inv=inventory(api,monkeypatch);req=inv['output']['items'][0]
    detail=api.get('/api/ingestion/documents/1',headers=headers()).json()
    approved=api.post('/api/ingestion/documents/1/review',headers=headers('admin'),json={'confirmed':True,'sha256':detail['sha256']})
    assert approved.status_code==200
    def decision(run,key,verification='source_confirmed'):
        return api.post(f'/api/applications/1/sources/runs/{run["id"]}/review',headers=headers('admin'),json={'item_key':key,'decision':'accepted','notes':'SYNTHETIC fixture review; not a real portal confirmation','source_url':'https://portal.example.test/fields','verification':verification,'output_hash':run['output_hash']})
    assert decision(inv,req['id'],'portal_confirmed').status_code==201
    monkeypatch.setattr(service.AIService,'chat',AsyncMock(return_value='Approved synthetic answer [source 1, paragraph 1].'))
    new=api.post('/api/applications/1/sources/draft',headers=headers(),json={'inventory_id':inv['id'],'requirement_id':req['id']}).json()
    asyncio.run(service.process(new['id']))
    run=next(r for r in api.get('/api/applications/1/sources',headers=headers()).json()['runs'] if r['id']==new['id'])
    assert decision(run,'draft').status_code==201
    assert decision(inv,'inventory_complete').status_code==201
    r=api.get('/api/applications/1/sources/export',headers=headers())
    with zipfile.ZipFile(io.BytesIO(r.content)) as archive:
        manifest=json.loads(archive.read('manifest.json'))
        assert manifest['package_status']=='REVIEWED EXPORT' and not manifest['submission_performed']
        assert any(f['path'].endswith('.docx') for f in manifest['files'])

@pytest.mark.parametrize('honda,expected',[(False,201),(True,422)])
def test_nonportal_official_forms_supported_but_honda_stays_portal_gated(api,monkeypatch,honda,expected):
    from routers.ingestion import router
    api.app.include_router(router)
    if not honda:
        async def clear_marker():
            async with database.AsyncSessionLocal() as db:
                app=await db.get(GrantApplication,1);app.form_data={};await db.commit()
        asyncio.run(clear_marker())
    inv=inventory(api,monkeypatch,{'1':'official_guidance'});req=inv['output']['items'][0]
    detail=api.get('/api/ingestion/documents/1',headers=headers()).json()
    assert api.post('/api/ingestion/documents/1/review',headers=headers('admin'),json={'confirmed':True,'sha256':detail['sha256']}).status_code==200
    data={'item_key':req['id'],'decision':'accepted','notes':'Synthetic complete official form fixture, not a portal','source_url':'https://funder.example.test/full-form.pdf','output_hash':inv['output_hash'],'inventory_basis':'official_complete_form'}
    assert api.post(f'/api/applications/1/sources/runs/{inv["id"]}/review',headers=headers('admin'),json=data).status_code==201
    data['item_key']='inventory_complete'
    assert api.post(f'/api/applications/1/sources/runs/{inv["id"]}/review',headers=headers('admin'),json=data).status_code==expected

def test_funder_instruction_cannot_match_itself_as_organization_evidence(api,monkeypatch):
    inv=inventory(api,monkeypatch,{'1':'official_guidance'})
    r=api.post('/api/applications/1/sources/match',headers=headers(),json={'inventory_id':inv['id'],'evidence_ids':[1]})
    assert r.status_code==202;asyncio.run(service.process(r.json()['id']))
    run=api.get('/api/applications/1/sources',headers=headers()).json()['runs'][0]
    assert run['status']=='completed' and run['output']['matches'][0]['evidence']==[]
    assert run['output']['matches'][0]['status']=='missing'

def test_existing_application_acceptance_does_not_bypass_new_source_requirements(api,monkeypatch):
    from models.models import Review
    from services.application_readiness import draft_revision
    async def accepted_before_analysis():
        async with database.AsyncSessionLocal() as db:
            app=await db.get(GrantApplication,1);app.proposal_text='Saved human draft'
            db.add(Review(application_id=1,reviewer_id=1,status='accepted',criteria={'application_revision':draft_revision(app)},feedback='Synthetic earlier review'))
            await db.commit()
    asyncio.run(accepted_before_analysis())
    inventory(api,monkeypatch)
    assert api.patch('/api/applications/1',json={'status':'submitted'},headers=headers('admin')).status_code==409
    assert api.post('/api/applications/1/reviews',json={'decision':'accepted','feedback':'Cannot bypass incomplete source workbench'},headers=headers('admin')).status_code==409
    assert api.get('/api/applications/1',headers=headers()).json()['status']=='draft'

def test_ai_queue_failure_does_not_fail_successful_extraction(api,monkeypatch):
    from config import settings
    from fastapi import HTTPException
    from models.models import AutomationJob,AutomationJobStatus
    from services import document_ingestion
    from sqlalchemy import select
    monkeypatch.setattr(settings,'APPLICATION_SOURCE_AUTO_ANALYZE',True)
    monkeypatch.setattr(service,'auto_enqueue',AsyncMock(side_effect=HTTPException(429,'Queue full')))
    async def run():
        async with database.AsyncSessionLocal() as db:
            item=await db.get(IngestionDocument,1);item.category='financials';await document_ingestion.queue(db,item,1);await db.commit()
            job=(await db.execute(select(AutomationJob).where(AutomationJob.job_type=='ingest'))).scalars().first();id=job.id
        await document_ingestion.process(id)
        async with database.AsyncSessionLocal() as db:
            item=await db.get(IngestionDocument,1);job=await db.get(AutomationJob,id)
            assert item.status=='needs_review' and job.status==AutomationJobStatus.completed
            assert item.error.startswith('Extraction succeeded;') and item.history[-1]['action']=='analysis_enqueue_failed'
    asyncio.run(run())
